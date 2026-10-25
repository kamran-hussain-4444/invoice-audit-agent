"""Write the current-year and post-year-end purchase ledgers (PLAN 15.5).

Sample invoices are posted as their spec says (including the planted errors). About 100 filler lines in
the current ledger and 50 in the post-YE ledger make the listings look like a real client export.
"""

import random
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from invoice_audit.money import quantize
from invoice_audit.synthetic.entities import CATEGORY_CODE, FILLER_SUPPLIERS, NOMINAL_ACCOUNTS, SUPPLIERS
from invoice_audit.synthetic.scenarios import FY_END, FY_START, POST_YE_END, InvoiceSpec, Scenario, ledger_code
from invoice_audit.synthetic.stable_zip import stable_zip

HEADERS = ["Posting Date", "Supplier", "Invoice No", "Description", "Nominal Code", "Nominal Account", "Net", "VAT",
           "Gross", "Currency", "Original Amount"]  # fmt: skip
FIXED_TIME = datetime(2026, 1, 1, 9, 0, 0)

FILLER_DESCRIPTIONS = {
    "Cost of sales": ["Raw materials", "Fixings and fasteners", "Packaging materials", "Carriage inwards",
                      "Castings", "Steel tube"],
    "Marketing": ["Signage", "Website banner design", "Promotional items"],
    "Professional fees": ["Payroll bureau fee", "HR advice"],
    "Utilities": ["Water rates", "Waste collection"],
    "IT & software": ["Telephone and broadband", "Mobile phone contract"],
    "Administrative expenses": ["Cleaning services", "Courier charges", "Refreshments", "Postage"],
    "Repairs & maintenance": ["Hydraulic hose replacement", "Safety equipment servicing", "Forklift service"],
    "Travel & subsistence": ["Hotel accommodation", "Staff travel"],
    "Premises (rent & rates)": ["Business rates instalment"],
    "Insurance": ["Insurance premium adjustment"],
    "Fixed assets (capital)": ["Workshop equipment"],
}  # fmt: skip


def _row(posting: date, supplier: str, number: str, description: str, code: str, net: Decimal, vat: Decimal,
         currency: str = "", original: Decimal | None = None) -> list:  # fmt: skip
    return [posting, supplier, number, description, code, NOMINAL_ACCOUNTS[code], net, vat, net + vat, currency,
            original]  # fmt: skip


def _sample_row(spec: InvoiceSpec, posting: date) -> list:
    number = spec.ledger_number or spec.invoice_number
    description = spec.description[:60]
    code = ledger_code(spec)
    if spec.currency != "GBP":
        original = spec.ledger_original if spec.ledger_original is not None else spec.net
        return _row(posting, spec.supplier.name, number, description, code, spec.ledger_gbp_net, Decimal("0.00"),
                    spec.currency, original)  # fmt: skip
    net = spec.ledger_net if spec.ledger_net is not None else spec.net
    return _row(posting, spec.supplier.name, number, description, code, net, spec.vat)


def _filler(scenario: Scenario, start: date, end: date, count: int) -> list[list]:
    rng: random.Random = scenario.rng
    rows = []
    known = list(SUPPLIERS.values())
    for i in range(count):
        posting = start + timedelta(days=rng.randint(0, (end - start).days))
        if rng.random() < 0.55:
            supplier = rng.choice(known)
            name, category = supplier.name, supplier.category
            number = scenario.numbering.number(supplier, posting)
            vat_rate, currency = supplier.vat_rate, supplier.currency
            low, high = supplier.amount_range
        else:
            name, category = FILLER_SUPPLIERS[i % len(FILLER_SUPPLIERS)]
            initials = "".join(word[0] for word in name.split()[:2])
            number = f"{initials}{rng.randint(10000, 99999)}"
            while number in scenario.numbering.used:
                number = f"{initials}{rng.randint(10000, 99999)}"
            scenario.numbering.used.add(number)
            vat_rate, currency = Decimal("20"), "GBP"
            low, high = 40, 2500
        net = Decimal(rng.randint(low, high)) + Decimal(rng.randint(0, 99)) / 100
        description = rng.choice(FILLER_DESCRIPTIONS[category])
        code = CATEGORY_CODE[category]
        if currency != "GBP":
            gbp = quantize(net * Decimal(rng.choice(["0.79", "0.85", "0.86"])))
            rows.append(_row(posting, name, number, description, code, gbp, Decimal("0.00"), currency, net))
        else:
            rows.append(_row(posting, name, number, description, code, net, quantize(net * vat_rate / 100)))
    return rows


def _write(path: Path, rows: list[list]) -> None:
    rows.sort(key=lambda r: (r[0], r[1], r[2]))
    wb = Workbook()
    wb.properties.creator = "Synthetic data generator"
    wb.properties.created = wb.properties.modified = FIXED_TIME
    ws = wb.active
    ws.title = "Purchase ledger"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for row in rows:
        ws.append(row)
    for row_cells in ws.iter_rows(min_row=2):
        row_cells[0].number_format = "DD/MM/YYYY"
        for cell in (row_cells[6], row_cells[7], row_cells[8], row_cells[10]):
            cell.number_format = "#,##0.00"
    for column, width in zip("ABCDEFGHIJK", (12, 34, 16, 44, 12, 26, 12, 11, 12, 9, 15), strict=True):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    buffer = BytesIO()
    wb.save(buffer)
    path.write_bytes(stable_zip(buffer.getvalue()))


def write_ledgers(scenario: Scenario, out: Path) -> tuple[int, int]:
    """Writes ledger_fy2026.xlsx and ledger_post_ye.xlsx. Returns the number of lines in each."""
    current: list[list] = []
    post_ye: list[list] = []
    for spec in [*scenario.invoices, scenario.blurry]:
        if spec.recorded_in in ("current", "both"):
            current.append(_sample_row(spec, spec.posting_date))
        if spec.recorded_in == "post_ye":
            post_ye.append(_sample_row(spec, spec.posting_date))
        if spec.recorded_in == "both":
            post_ye.append(_sample_row(spec, spec.post_ye_posting_date))
    current += _filler(scenario, FY_START, FY_END, 100)
    post_ye += _filler(scenario, FY_END + timedelta(days=1), POST_YE_END, 50)
    _write(out / "ledger_fy2026.xlsx", current)
    _write(out / "ledger_post_ye.xlsx", post_ye)
    return len(current), len(post_ye)
