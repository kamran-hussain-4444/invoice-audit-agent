"""Excel invoices laid out like a typical small-business invoice template."""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from invoice_audit.synthetic.view import InvoiceView

FIXED_TIME = datetime(2026, 1, 1, 9, 0, 0)


def render_invoice_xlsx(view: InvoiceView) -> bytes:
    wb = Workbook()
    wb.properties.creator = view.supplier_name
    wb.properties.created = wb.properties.modified = FIXED_TIME
    ws = wb.active
    ws.title = "Invoice"
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 18

    ws["A1"] = view.supplier_name
    ws["A1"].font = Font(bold=True, size=16)
    ws["B1"] = view.title
    ws["B1"].font = Font(bold=True, size=14)
    row = 2
    for line in view.supplier_lines:
        ws.cell(row, 1, line).font = Font(size=9)
        row += 1
    start = row + 1
    ws.cell(start, 1, "Bill to").font = Font(bold=True)
    for i, line in enumerate(view.bill_to):
        ws.cell(start + 1 + i, 1, line)
    for i, (label, value) in enumerate(view.meta):
        ws.cell(start + 1 + i, 2, f"{label}: {value}")
    row = start + 2 + max(len(view.bill_to), len(view.meta))
    if view.note:
        ws.cell(row, 1, view.note).font = Font(italic=True)
        row += 2

    thin = Side(style="thin", color="999999")
    fill = PatternFill("solid", fgColor="DDEBF7")
    for col, text in ((1, "Description"), (2, f"Amount ({view.currency})")):
        cell = ws.cell(row, col, text)
        cell.font, cell.fill, cell.border = Font(bold=True), fill, Border(bottom=thin)
    row += 1
    for text, amount in view.lines:
        ws.cell(row, 1, text)
        ws.cell(row, 2, amount).alignment = Alignment(horizontal="right")
        row += 1
    row += 1
    for label, value in view.totals:
        ws.cell(row, 1, label).alignment = Alignment(horizontal="right")
        total = ws.cell(row, 2, value)
        total.alignment = Alignment(horizontal="right")
        total.font = Font(bold=label.startswith("Total"))
        row += 1
    ws.cell(row + 1, 1, view.footer).font = Font(size=8, italic=True)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
