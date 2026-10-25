"""The Excel workpaper (PLAN section 14) and run_results.json.

Six tabs: Summary, Extracted Data, Main Testing, Cut-off Testing, Exceptions, Processing Log.
Neutral styling, no firm branding. Prepared/reviewed sign-off cells are left blank for the auditor.
"""

import json
import re
import shutil
from collections import Counter
from dataclasses import asdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from invoice_audit.extraction.prompt import PROMPT_VERSION
from invoice_audit.models import CheckResult, CutoffTestRow, ExtractionRecord, MainTestRow
from invoice_audit.procedures.cutoff import RECORDED_LABELS
from invoice_audit.procedures.main_testing import value_tested
from invoice_audit.validation import CONFIDENCE_FIELDS

HEADER_FILL = PatternFill("solid", fgColor="2F3E4E")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FILLS = {
    "pass": PatternFill("solid", fgColor="C6EFCE"),
    "exception": PatternFill("solid", fgColor="FFC7CE"),
    "review": PatternFill("solid", fgColor="FFEB9C"),
}
AMBER = FILLS["review"]
STATUS_FILLS = {"ok": FILLS["pass"], "needs_review": AMBER, "excluded": AMBER, "failed": FILLS["exception"]}
TICKS = {"pass": "✓", "exception": "✗", "review": "?", "n/a": "–"}
TICK_FONT = Font(name="Segoe UI Symbol", size=11)
MONEY = "#,##0.00"
DATE = "DD/MM/YYYY"
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
RESOLUTIONS = ["Open", "Cleared", "Adjusted", "Passed to ML"]
DISCLAIMER = ("Data extracted with AI assistance. Exceptions and review items require auditor judgement. "
              "This schedule does not replace professional scepticism.")  # fmt: skip
POSITION_LABELS = {"pre_ye": "Pre-YE", "post_ye": "Post-YE", "spans": "Spans YE"}


def display_path(path: Path | None) -> str:
    """Paths relative to the project folder where possible (no user names in a shared workpaper)."""
    if path is None:
        return ""
    try:
        return Path(path).resolve().relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return Path(path).name


def company_slug(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_")


# ---------------------------------------------------------------- generic table writer


def _table(ws: Worksheet, headers: list[str], rows: list[list], widths: dict[str, int], start_row: int = 1) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(start_row, col, header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    money_cols, date_cols = set(), set()
    for r, values in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(r, c, value)
            cell.alignment = TOP
            if isinstance(value, Decimal):
                money_cols.add(c)
            elif isinstance(value, date):
                date_cols.add(c)
    # Format whole columns, so blank cells an auditor fills in later get the right format too.
    for c in money_cols | date_cols:
        for r in range(start_row + 1, start_row + 1 + len(rows)):
            ws.cell(r, c).number_format = MONEY if c in money_cols else DATE
    for col, header in enumerate(headers, start=1):
        width = widths.get(header, 14)
        ws.column_dimensions[get_column_letter(col)].width = width
        if width >= 30:
            for r in range(start_row + 1, start_row + 1 + len(rows)):
                ws.cell(r, col).alignment = WRAP
    last = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{start_row}:{last}{start_row + max(len(rows), 1)}"
    ws.freeze_panes = ws.cell(start_row + 1, 2)
    _print_setup(ws)


def _print_setup(ws: Worksheet) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _tick(ws: Worksheet, row: int, col: int, check: CheckResult) -> None:
    cell = ws.cell(row, col, TICKS[check.status])
    cell.font = TICK_FONT
    cell.alignment = Alignment(horizontal="center", vertical="top")
    if check.status in FILLS:
        cell.fill = FILLS[check.status]


def _result_cell(ws: Worksheet, row: int, col: int, status: str) -> None:
    cell = ws.cell(row, col, status.capitalize())
    cell.fill = FILLS[status]


def _notes(checks: dict[str, CheckResult]) -> str:
    return "; ".join(f"{name}: {c.detail}" for name, c in checks.items() if c.status in ("exception", "review"))


# ---------------------------------------------------------------- tabs


def _summary(ws: Worksheet, result) -> None:
    cfg, stats = result.cfg, result.stats
    e, i = cfg.engagement, cfg.inputs
    records = result.records
    status_counts = Counter(r.status for r in records)
    findings_by_test = Counter((f.test, f.status) for f in result.findings)

    rows: list[tuple] = [
        ("Purchases: tests of detail and cut-off",),
        (),
        ("Client", e.company_name),
        ("Period", f"{e.fy_start:%d/%m/%Y} to {e.fy_end:%d/%m/%Y}"),
        ("Workpaper", "Purchases: tests of detail and cut-off"),
        ("Prepared by", "", "Date", ""),
        ("Reviewed by", "", "Date", ""),
        (),
        ("Objectives",),
        ("Test 1: Main testing", "To confirm that sampled purchase invoices are recorded at the correct amount, in "
         "the correct period and nominal account, and relate to the audited entity (accuracy, occurrence, "
         "classification)."),
        ("Test 2: Cut-off", "To confirm that purchases either side of the year end are recorded in the correct "
         "period (cut-off), identifying any accruals or prepayments needed."),
        (),
        ("Settings used",),
        ("Main sample folder", display_path(i.main_folder)),
        ("Cut-off sample folder", display_path(i.cutoff_folder)),
        ("Current-year ledger", display_path(i.current_ledger)),
        ("Post-year-end ledger", display_path(i.post_ye_ledger)),
        ("Ledger type / basis", f"{i.ledger_type} (compare {cfg.amount_basis})"),
        ("Value tolerance (£)", cfg.matching.value_tolerance),
        ("Model / effort", f"{cfg.ai.model} / {cfg.ai.effort}"),
        ("Prompt version", PROMPT_VERSION),
        (),
        ("Document counts",),
        ("Files found", stats.files_found),
        ("Documents extracted", sum(1 for r in records if r.invoice is not None)),
        ("From cache", stats.cache_hits),
        ("Needs review", status_counts["needs_review"]),
        ("Excluded", status_counts["excluded"]),
        ("Failed", status_counts["failed"]),
        (),
    ]  # fmt: skip
    for row_index, values in enumerate(rows, start=1):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_index, col, value)
            if isinstance(value, Decimal):
                cell.number_format = MONEY
    ws["A1"].font = Font(bold=True, size=14)
    for label in ("Objectives", "Settings used", "Document counts"):
        next(c for c in ws["A"] if c.value == label).font = Font(bold=True)
    for r in (6, 7):
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 3).font = Font(bold=True)
    for r in (10, 11):  # the two objective paragraphs
        ws.cell(r, 2).alignment = WRAP
        ws.cell(r, 1).alignment = TOP
        ws.row_dimensions[r].height = 48

    # Per-test results table.
    row = len(rows) + 1
    ws.cell(row, 1, "Results by test").font = Font(bold=True)
    row += 1
    headers = ["Test", "Items tested", "Pass", "Exception", "Review", "Value tested (£)", "Value of exceptions (£)"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row, col, header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for name, test_rows, enabled in (("Main", result.main_rows, "main" in cfg.tests.run),
                                     ("Cut-off", result.cutoff_rows, "cutoff" in cfg.tests.run)):  # fmt: skip
        if not enabled:
            continue
        row += 1
        counts = Counter(r.overall for r in test_rows)
        tested = value_tested(test_rows, cfg)
        impact = sum((f.amount_impact or Decimal(0) for f in result.findings
                      if f.test == name and f.status == "exception"), Decimal("0.00"))  # fmt: skip
        values = [name, len(test_rows), counts["pass"], counts["exception"], counts["review"], tested, impact]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            if isinstance(value, Decimal):
                cell.number_format = MONEY
        ws.cell(row, 3).fill, ws.cell(row, 4).fill, ws.cell(row, 5).fill = (
            FILLS["pass"], FILLS["exception"], FILLS["review"])  # fmt: skip
    row += 1
    extraction = findings_by_test[("Extraction", "review")]
    if extraction:
        row += 1
        ws.cell(row, 1, "Extraction findings (review)")
        ws.cell(row, 2, extraction)

    # Run stats.
    t = stats.tokens
    run_rows = [
        ("Run statistics",),
        ("Mode", stats.mode),
        ("API calls", stats.live_calls + stats.batch_calls),
        ("Cache hits", stats.cache_hits),
        ("Input / output / cache-read tokens", f"{t.input_tokens:,} / {t.output_tokens:,} / {t.cache_read_tokens:,}"),
        ("Estimated cost (USD)", float(round(stats.cost_usd, 4))),
        ("Duration (seconds)", round(stats.duration_seconds, 1)),
        ("Documents per minute", round(stats.docs_per_minute, 1)),
        (),
        ("Tick marks",),
        ("✓", "Pass: agreed without exception"),
        ("✗", "Exception: see the Exceptions tab"),
        ("?", "Review: auditor judgement needed (see the Exceptions tab)"),
        ("–", "Not applicable / not tested"),
        (),
        (DISCLAIMER,),
    ]
    start = row + 2
    for offset, values in enumerate(run_rows):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(start + offset, col, value)
            if col == 1 and value in TICKS.values():
                cell.font = TICK_FONT
    for label in ("Run statistics", "Tick marks"):
        next(c for c in ws["A"] if c.value == label).font = Font(bold=True)
    ws.cell(start + len(run_rows) - 1, 1).font = Font(italic=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 70
    for col in "CDEFG":
        ws.column_dimensions[col].width = 16
    _print_setup(ws)


EXTRACTED_HEADERS = [
    "Doc ref", "File", "Test", "Status", "Reasons", "Document type", "Supplier", "Supplier VAT no", "Addressed to",
    "Invoice no", "Invoice date", "Currency", "Net", "VAT", "Gross", "Description", "Delivery date",
    "Service period start", "Service period end", "Period basis", "Period reasoning", "Suggested category",
    "Category reasoning", *[f"Conf: {k.replace('_', ' ')}" for k in CONFIDENCE_FIELDS], "Extraction notes", "Source",
]  # fmt: skip
# Which columns turn amber when a field is weak.
WEAK_COLUMNS = {
    "supplier_name": "Supplier", "addressed_to": "Addressed to", "invoice_number": "Invoice no",
    "invoice_date": "Invoice date", "currency": "Currency", "net_amount": "Net", "vat_amount": "VAT",
    "gross_amount": "Gross", "description": "Description", "service_period": "Service period start",
}  # fmt: skip


def _extracted(ws: Worksheet, records: list[ExtractionRecord]) -> None:
    rows = []
    for n, record in enumerate(records, start=1):
        inv = record.invoice
        base = [f"DOC-{n:03d}", record.source.display_name, "Main" if record.source.test == "main" else "Cut-off",
                record.status.replace("_", " ").capitalize(), "; ".join(r for r in record.reasons if r)]  # fmt: skip
        if inv is None:
            rows.append(base + [""] * (len(EXTRACTED_HEADERS) - len(base) - 1) + [record.mode or ""])
            continue
        confidence = [getattr(inv.field_confidence, k) for k in CONFIDENCE_FIELDS]
        rows.append(base + [
            inv.document_type.replace("_", " "), inv.supplier_name, inv.supplier_vat_number, inv.addressed_to,
            inv.invoice_number, inv.invoice_date, inv.currency, inv.net_amount, inv.vat_amount, inv.gross_amount,
            inv.description, inv.delivery_date, inv.service_period_start, inv.service_period_end,
            inv.period_basis.replace("_", " "), inv.period_reasoning, inv.suggested_category, inv.category_reasoning,
            *confidence, inv.extraction_notes, record.mode or "",
        ])  # fmt: skip
    widths = {"File": 30, "Reasons": 34, "Supplier": 26, "Addressed to": 26, "Description": 40,
              "Period reasoning": 40, "Category reasoning": 40, "Extraction notes": 34,
              "Suggested category": 20}  # fmt: skip
    _table(ws, EXTRACTED_HEADERS, rows, widths)
    column = {h: i for i, h in enumerate(EXTRACTED_HEADERS, start=1)}
    for r, record in enumerate(records, start=2):
        ws.cell(r, column["Status"]).fill = STATUS_FILLS[record.status]
        if record.invoice is None:
            continue
        for field in record.weak_fields:
            if field in WEAK_COLUMNS:
                ws.cell(r, column[WEAK_COLUMNS[field]]).fill = AMBER
        for k in CONFIDENCE_FIELDS:
            cell = ws.cell(r, column[f"Conf: {k.replace('_', ' ')}"])
            if cell.value == "low":
                cell.fill = AMBER


MAIN_HEADERS = ["Ref", "File", "Supplier", "Invoice no", "Invoice date", "Ledger row", "Posting date", "Nominal code",
                "Nominal account", "Ledger amount", "Invoice amount (basis)", "Difference", "A", "D", "E", "C",
                "Suggested category", "Overall result", "Finding refs", "Notes"]  # fmt: skip


def _main_testing(ws: Worksheet, rows: list[MainTestRow]) -> None:
    values = []
    for row in rows:
        inv, line = row.record.invoice, row.ledger_line
        values.append([
            row.ref, row.record.source.display_name, inv.supplier_name, inv.invoice_number, inv.invoice_date,
            line.row_number if line else None, line.posting_date if line else None, line.nominal_code if line else None,
            line.nominal_account if line else None, row.ledger_amount, row.invoice_amount, row.difference,
            None, None, None, None, inv.suggested_category, None, ", ".join(row.finding_refs),
            "; ".join(filter(None, [row.match.detail if row.match.status != "pass" else "",
                                    _notes({"A": row.A, "D": row.D, "E": row.E, "C": row.C})])),
        ])  # fmt: skip
    widths = {"File": 30, "Supplier": 26, "Nominal account": 22, "Ledger amount": 14, "Invoice amount (basis)": 14,
              "A": 5, "D": 5, "E": 5, "C": 5, "Suggested category": 20, "Notes": 46, "Finding refs": 16}  # fmt: skip
    _table(ws, MAIN_HEADERS, values, widths)
    column = {h: i for i, h in enumerate(MAIN_HEADERS, start=1)}
    for r, row in enumerate(rows, start=2):
        for name in "ADEC":
            _tick(ws, r, column[name], getattr(row, name))
        _result_cell(ws, r, column["Overall result"], row.overall)


CUTOFF_HEADERS = ["Ref", "File", "Supplier", "Invoice no", "Invoice date", "Period basis", "Period start",
                  "Period end", "Position", "Recorded in", "Posting date", "Ledger amount", "A", "P",
                  "Apportionment working", "Suggested adjustment (£)", "Overall result", "Finding refs",
                  "Notes"]  # fmt: skip


def _cutoff(ws: Worksheet, rows: list[CutoffTestRow]) -> None:
    values = []
    for row in rows:
        inv, line = row.record.invoice, row.ledger_line
        values.append([
            row.ref, row.record.source.display_name, inv.supplier_name, inv.invoice_number, inv.invoice_date,
            row.period_basis, row.period_start, row.period_end, POSITION_LABELS.get(row.position, ""),
            RECORDED_LABELS.get(row.recorded_in, ""), line.posting_date if line else None, row.ledger_amount,
            None, None, row.apportionment, row.suggested_adjustment, None, ", ".join(row.finding_refs),
            "; ".join(filter(None, [row.match.detail if row.match.status != "pass" else "",
                                    _notes({"A": row.A, "P": row.P})])),
        ])  # fmt: skip
    widths = {"File": 30, "Supplier": 26, "Period basis": 18, "A": 5, "P": 5, "Apportionment working": 32,
              "Suggested adjustment (£)": 14, "Notes": 46, "Finding refs": 16}  # fmt: skip
    _table(ws, CUTOFF_HEADERS, values, widths)
    column = {h: i for i, h in enumerate(CUTOFF_HEADERS, start=1)}
    for r, row in enumerate(rows, start=2):
        _tick(ws, r, column["A"], row.A)
        _tick(ws, r, column["P"], row.P)
        _result_cell(ws, r, column["Overall result"], row.overall)


EXCEPTION_HEADERS = ["Ref", "Test", "Test ref", "Status", "Issue type", "File", "Supplier", "Invoice no",
                     "Explanation", "AI reasoning", "Suggested action", "Amount impact (£)", "Auditor conclusion",
                     "Resolution"]  # fmt: skip


def _exceptions(ws: Worksheet, findings) -> None:
    values = [[f.ref, f.test, f.test_ref, f.status.capitalize(), f.issue_type, f.file, f.supplier, f.invoice_number,
               f.explanation, f.ai_reasoning, f.suggested_action, f.amount_impact, None, "Open"]
              for f in findings]  # fmt: skip
    widths = {"Issue type": 22, "File": 28, "Supplier": 24, "Explanation": 60, "AI reasoning": 40,
              "Suggested action": 36, "Amount impact (£)": 14, "Auditor conclusion": 34, "Resolution": 14}  # fmt: skip
    _table(ws, EXCEPTION_HEADERS, values, widths)
    column = {h: i for i, h in enumerate(EXCEPTION_HEADERS, start=1)}
    for r, finding in enumerate(findings, start=2):
        ws.cell(r, column["Status"]).fill = FILLS[finding.status]
    letter = get_column_letter(column["Resolution"])
    validation = DataValidation(type="list", formula1='"' + ",".join(RESOLUTIONS) + '"', allow_blank=True)
    validation.error, validation.errorTitle = "Choose Open, Cleared, Adjusted or Passed to ML", "Resolution"
    ws.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{max(len(findings), 1) + 200}")


def _processing_log(ws: Worksheet, records: list[ExtractionRecord], log_path: Path | None) -> None:
    problems = [
        r for r in records if r.status in ("failed", "excluded") or (r.status == "needs_review" and not r.invoice)
    ]
    values = [[r.source.display_name, "Main" if r.source.test == "main" else "Cut-off",
               r.status.replace("_", " ").capitalize(), "; ".join(r.reasons)] for r in problems]  # fmt: skip
    _table(ws, ["File", "Test", "Status", "Reason"], values, {"File": 40, "Reason": 70})
    for r, record in enumerate(problems, start=2):
        ws.cell(r, 3).fill = FILLS["exception"] if record.status == "failed" else AMBER
    row = len(problems) + 3
    if not problems:
        ws.cell(2, 1, "No documents failed or were excluded.")
    ws.cell(row, 1, "Detailed log file").font = Font(bold=True)
    ws.cell(row, 2, display_path(log_path) if log_path else "(no log file)")


# ---------------------------------------------------------------- entry points


def write_workpaper(result, folder: Path) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    cfg = result.cfg
    name = f"{company_slug(cfg.engagement.company_name)}_FY{cfg.engagement.fy_end:%Y}_invoice_testing_" \
           f"{result.started_at:%Y%m%d_%H%M%S}.xlsx"  # fmt: skip
    wb = Workbook()
    wb.properties.creator = "invoice-audit-agent"
    sheets = [("Summary", lambda ws: _summary(ws, result)),
              ("Extracted Data", lambda ws: _extracted(ws, result.records)),
              ("Main Testing", lambda ws: _main_testing(ws, result.main_rows)),
              ("Cut-off Testing", lambda ws: _cutoff(ws, result.cutoff_rows)),
              ("Exceptions", lambda ws: _exceptions(ws, result.findings)),
              ("Processing Log", lambda ws: _processing_log(ws, result.records, result.log_path))]  # fmt: skip
    for index, (title, fill) in enumerate(sheets):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        fill(ws)
    path = folder / name
    wb.save(path)
    return path


def _jsonable(value):
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Path):
        return display_path(value)
    if isinstance(value, set):
        return sorted(value)
    if isinstance(value, bytes):
        return None
    raise TypeError(f"Not JSON serialisable: {type(value)}")


def _check(check: CheckResult) -> dict:
    return {"status": check.status, "detail": check.detail, "finding_ref": check.finding_ref}


def write_results_json(result, folder: Path) -> Path:
    """output/run_results.json (full detail) and output/latest_run.json (same file, for Power Automate)."""
    stats = result.stats
    exceptions = sum(1 for f in result.findings if f.status == "exception")
    data = {
        "workpaper": display_path(result.workpaper_path) if result.workpaper_path else None,
        "workpaper_absolute": str(result.workpaper_path.resolve()) if result.workpaper_path else None,
        "company": result.cfg.engagement.company_name,
        "started_at": result.started_at,
        "counts": {
            "documents": stats.documents,
            "failed": sum(1 for r in result.records if r.status == "failed"),
            "excluded": sum(1 for r in result.records if r.status == "excluded"),
            "needs_review": sum(1 for r in result.records if r.status == "needs_review"),
            "exceptions": exceptions,
            "reviews": sum(1 for f in result.findings if f.status == "review"),
            "main_tested": len(result.main_rows),
            "cutoff_tested": len(result.cutoff_rows),
        },
        "stats": {**asdict(stats), "mode": stats.mode, "docs_per_minute": stats.docs_per_minute},
        "records": [
            {
                "file": r.source.display_name,
                "test": r.source.test,
                "status": r.status,
                "reasons": r.reasons,
                "weak_fields": r.weak_fields,
                "from_cache": r.from_cache,
                "invoice": r.invoice.model_dump() if r.invoice else None,
            }
            for r in result.records
        ],
        "main_testing": [
            {
                "ref": row.ref,
                "file": row.record.source.display_name,
                "overall": row.overall,
                "checks": {k: _check(getattr(row, k)) for k in ("match", "A", "D", "E", "C")},
                "ledger_amount": row.ledger_amount,
                "invoice_amount": row.invoice_amount,
                "difference": row.difference,
            }
            for row in result.main_rows
        ],  # fmt: skip
        "cutoff_testing": [
            {
                "ref": row.ref,
                "file": row.record.source.display_name,
                "overall": row.overall,
                "position": row.position,
                "recorded_in": row.recorded_in,
                "period_start": row.period_start,
                "period_end": row.period_end,
                "suggested_adjustment": row.suggested_adjustment,
                "checks": {k: _check(getattr(row, k)) for k in ("match", "A", "P")},
            }
            for row in result.cutoff_rows
        ],  # fmt: skip
        "findings": [asdict(f) for f in result.findings],
    }
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / "run_results.json"
    path.write_text(json.dumps(data, indent=2, default=_jsonable, ensure_ascii=False), encoding="utf-8")
    shutil.copyfile(path, folder / "latest_run.json")
    return path
