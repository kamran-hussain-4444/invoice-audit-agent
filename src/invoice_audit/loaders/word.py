"""Word (.docx) invoices converted to text locally: header, body and tables in document order, footer.

Suppliers often put their name and VAT number in the page header, so headers must not be skipped.
"""

from io import BytesIO

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from invoice_audit.loaders.errors import LoaderError


def _table_text(table: Table) -> list[str]:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = " ".join(cell.text.split())
            if not cells or text != cells[-1]:  # merged cells repeat their text
                cells.append(text)
        if any(cells):
            rows.append("\t".join(cells))
    return rows


def docx_text(data: bytes) -> str:
    try:
        doc = Document(BytesIO(data))
    except Exception:
        raise LoaderError("Corrupt or unreadable Word document") from None

    header, footer = [], []
    for section in doc.sections:
        for part, target in ((section.header, header), (section.footer, footer)):
            for paragraph in part.paragraphs:
                if paragraph.text.strip() and paragraph.text not in target:
                    target.append(paragraph.text.strip())
            for table in part.tables:
                target.extend(_table_text(table))

    body = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, doc).text.strip()
            if text:
                body.append(text)
        elif tag == "tbl":
            body.extend(_table_text(Table(child, doc)))

    if not (header or body or footer):
        raise LoaderError("No text found in Word document")
    parts = []
    if header:
        parts += ["[Page header]", *header, ""]
    parts += ["[Document]", *body]
    if footer:
        parts += ["", "[Page footer]", *footer]
    return "\n".join(parts)
