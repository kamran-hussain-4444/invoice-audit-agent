"""Excel (.xlsx) invoices converted to a tab-separated grid per visible sheet."""

from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from invoice_audit.loaders.errors import LoaderError

MAX_ROWS_PER_SHEET = 500


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = value.date() if value.time() == datetime.min.time() else value
    if isinstance(value, date) and not isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float):
        return f"{value:.2f}" if value != int(value) else str(int(value))
    return " ".join(str(value).split())


def xlsx_text(data: bytes) -> str:
    try:
        workbook = load_workbook(BytesIO(data), data_only=True)  # data_only: formula results, not formulas
    except Exception:
        raise LoaderError("Corrupt or unreadable Excel workbook") from None

    parts = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            continue
        rows = []
        for values in sheet.iter_rows(values_only=True):
            cells = [_cell(v) for v in values]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                rows.append("\t".join(cells))
            if len(rows) >= MAX_ROWS_PER_SHEET:
                rows.append(f"... (more rows not shown, limit {MAX_ROWS_PER_SHEET})")
                break
        if rows:
            parts += [f"[Sheet: {sheet.title}]", *rows, ""]
    workbook.close()
    if not parts:
        raise LoaderError("No data found in Excel workbook")
    return "\n".join(parts).strip()
