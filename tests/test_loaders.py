"""Loaders and discovery. Small fixture files are built in a temp folder, so nothing binary is committed."""

import base64
from email import message_from_bytes, policy
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from reportlab.pdfgen.canvas import Canvas

from invoice_audit.discovery import discover
from invoice_audit.loaders import LoaderError, NeedsReviewWithoutExtraction, load
from invoice_audit.loaders.email import html_to_text, parse_eml
from invoice_audit.models import SourceDocument
from tests.conftest import ROOT

FIXTURES = ROOT / "tests" / "fixtures"


def make_pdf(pages: int = 1, text: str = "INVOICE INV-1") -> bytes:
    buffer = BytesIO()
    c = Canvas(buffer, invariant=1)
    for i in range(pages):
        c.drawString(72, 720, f"{text} page {i + 1}")
        c.showPage()
    c.save()
    return buffer.getvalue()


def big_photo(size=(3000, 2000), fmt="JPEG") -> bytes:
    """Noisy enough that it is well over the 20 KB logo threshold."""
    image = Image.effect_noise(size, 60).convert("RGB")
    out = BytesIO()
    image.save(out, format=fmt, quality=85)
    return out.getvalue()


def doc(path: Path, kind: str, **kw) -> SourceDocument:
    return SourceDocument(doc_id="x", test="main", file_path=path, display_name=path.name, kind=kind, **kw)


# ---------------------------------------------------------------- PDF


def test_pdf_sent_as_document_block(tmp_path, cfg):
    path = tmp_path / "invoice_07.pdf"
    path.write_bytes(make_pdf())
    payload = load(doc(path, "pdf"), cfg)
    assert payload.blocks[0] == {"type": "text", "text": "Source: invoice_07.pdf (PDF)"}
    block = payload.blocks[1]
    assert block["type"] == "document"
    assert block["source"]["media_type"] == "application/pdf"
    assert base64.b64decode(block["source"]["data"]) == path.read_bytes()
    assert payload.approx_bytes > 0


def test_truncated_pdf_fails(tmp_path, cfg):
    data = make_pdf()
    path = tmp_path / "bad.pdf"
    path.write_bytes(data[: len(data) // 2])
    with pytest.raises(LoaderError, match="Corrupt or unreadable PDF"):
        load(doc(path, "pdf"), cfg)


def test_not_a_pdf_fails(tmp_path, cfg):
    path = tmp_path / "fake.pdf"
    path.write_bytes(b"hello, I am a text file")
    with pytest.raises(LoaderError, match="Corrupt"):
        load(doc(path, "pdf"), cfg)


def test_encrypted_pdf_fails(tmp_path, cfg):
    writer = PdfWriter(clone_from=BytesIO(make_pdf()))
    writer.encrypt("secret")
    path = tmp_path / "locked.pdf"
    with path.open("wb") as f:
        writer.write(f)
    with pytest.raises(LoaderError, match="Password-protected PDF"):
        load(doc(path, "pdf"), cfg)


def test_long_pdf_goes_to_review_without_api_call(tmp_path, cfg):
    path = tmp_path / "long.pdf"
    path.write_bytes(make_pdf(pages=cfg.ai.max_pdf_pages + 1))
    with pytest.raises(NeedsReviewWithoutExtraction, match="21 pages"):
        load(doc(path, "pdf"), cfg)


# ---------------------------------------------------------------- images


def test_photo_resized_to_max_edge(tmp_path, cfg):
    path = tmp_path / "photo.jpg"
    path.write_bytes(big_photo())
    payload = load(doc(path, "image"), cfg)
    block = payload.blocks[1]
    assert block["type"] == "image"
    image = Image.open(BytesIO(base64.b64decode(block["source"]["data"])))
    assert max(image.size) == cfg.ai.image_max_edge


def test_exif_rotation_applied(tmp_path, cfg):
    image = Image.new("RGB", (400, 200), "white")
    exif = image.getexif()
    exif[0x0112] = 6  # "rotate 90° clockwise to display"
    path = tmp_path / "sideways.jpg"
    image.save(path, exif=exif)
    payload = load(doc(path, "image"), cfg)
    out = Image.open(BytesIO(base64.b64decode(payload.blocks[1]["source"]["data"])))
    assert out.size == (200, 400)


def test_multipage_tiff_gives_one_block_per_page(tmp_path, cfg):
    frames = [Image.new("RGB", (300, 400), colour) for colour in ("white", "grey", "black")]
    path = tmp_path / "scan.tiff"
    frames[0].save(path, save_all=True, append_images=frames[1:])
    payload = load(doc(path, "image"), cfg)
    assert [b["type"] for b in payload.blocks] == ["text", "image", "image", "image"]


def test_corrupt_image_fails(tmp_path, cfg):
    path = tmp_path / "broken.png"
    path.write_bytes(b"\x89PNG not really")
    with pytest.raises(LoaderError, match="image"):
        load(doc(path, "image"), cfg)


# ---------------------------------------------------------------- Word / Excel


def test_docx_header_body_table_in_order(tmp_path, cfg):
    d = Document()
    d.sections[0].header.paragraphs[0].text = "Acme Widgets Ltd, VAT GB123456789"
    d.add_paragraph("Invoice No: AW-1001")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Description", "Amount"
    table.cell(1, 0).text, table.cell(1, 1).text = "Widgets", "100.00"
    d.add_paragraph("Total due: £120.00")
    d.sections[0].footer.paragraphs[0].text = "Registered in England"
    path = tmp_path / "inv.docx"
    d.save(path)
    text = load(doc(path, "docx"), cfg).blocks[1]["text"]
    assert text.index("Acme Widgets Ltd") < text.index("AW-1001") < text.index("Widgets\t100.00")
    assert text.index("Widgets\t100.00") < text.index("Total due") < text.index("Registered in England")


def test_empty_docx_fails(tmp_path, cfg):
    path = tmp_path / "empty.docx"
    Document().save(path)
    with pytest.raises(LoaderError, match="No text found in Word document"):
        load(doc(path, "docx"), cfg)


def test_xlsx_grid_visible_sheets_only(tmp_path, cfg):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.append(["Supplier", "Acme Ltd"])
    ws.append(["Net", 100.5])
    ws.append([])
    ws.append(["Total", "=B2*1.2"])
    hidden = wb.create_sheet("Workings")
    hidden.append(["secret"])
    hidden.sheet_state = "hidden"
    path = tmp_path / "inv.xlsx"
    wb.save(path)
    text = load(doc(path, "xlsx"), cfg).blocks[1]["text"]
    assert "[Sheet: Invoice]" in text
    assert "Supplier\tAcme Ltd" in text
    assert "Net\t100.50" in text
    assert "secret" not in text


# ---------------------------------------------------------------- email


def make_eml(attachments=(), body="Please find attached.", html: str | None = None, subject="Invoice") -> bytes:
    message = EmailMessage()
    message["From"] = "Supplier <billing@supplier.example>"
    message["To"] = "ap@client.example"
    message["Subject"] = subject
    message.set_content(body)
    if html:
        message.add_alternative(html, subtype="html")
    for name, data, mime in attachments:
        main, sub = mime.split("/")
        message.add_attachment(data, maintype=main, subtype=sub, filename=name)
    return message.as_bytes()


def test_email_attachments_split_and_logo_skipped(tmp_path):
    logo = BytesIO()
    Image.new("RGB", (60, 20), "blue").save(logo, format="PNG")
    eml = make_eml([("logo.png", logo.getvalue(), "image/png"), ("INV-1.pdf", make_pdf(), "application/pdf"),
                    ("INV-2.pdf", make_pdf(text="INVOICE INV-2"), "application/pdf")])  # fmt: skip
    (tmp_path / "email_01.eml").write_bytes(eml)
    found = discover(tmp_path, "main")
    assert [d.display_name for d in found.documents] == ["email_01.eml › INV-1.pdf", "email_01.eml › INV-2.pdf"]
    assert all(d.kind == "pdf" and d.content_bytes for d in found.documents)
    assert "Subject: Invoice" in found.documents[0].email_context
    assert found.documents[0].doc_id != found.documents[1].doc_id


def test_email_context_block_added(tmp_path, cfg):
    (tmp_path / "e.eml").write_bytes(make_eml([("INV-1.pdf", make_pdf(), "application/pdf")]))
    (document,) = discover(tmp_path, "main").documents
    payload = load(document, cfg)
    assert payload.blocks[0]["text"] == "Source: e.eml › INV-1.pdf (PDF)"
    assert payload.blocks[1]["text"].startswith("Email context:\nSubject: Invoice")
    assert payload.blocks[2]["type"] == "document"


def test_email_without_attachment_uses_body(tmp_path, cfg):
    (tmp_path / "e.eml").write_bytes(make_eml(body="INVOICE 77\nTotal £50.00", subject="Invoice 77"))
    (document,) = discover(tmp_path, "main").documents
    assert document.kind == "email_body"
    assert document.display_name == "e.eml (email body)"
    text = load(document, cfg).blocks[1]["text"]
    assert "INVOICE 77" in text and "Subject: Invoice 77" in text


def test_html_only_body_is_stripped():
    eml = make_eml(
        body="", html="<html><body><p>Invoice&nbsp;99</p><table><tr><td>Total</td></tr></table></body></html>"
    )
    parsed = parse_eml(eml)
    assert "Invoice 99" in parsed.body
    assert "<" not in html_to_text("<p>a</p><br><b>b</b>")


def test_attached_email_expanded_one_level(tmp_path):
    raw = make_eml([("INV-9.pdf", make_pdf(), "application/pdf")], subject="Fwd inner")
    inner = message_from_bytes(raw, policy=policy.default)
    outer = EmailMessage()
    outer["Subject"] = "FW: invoice"
    outer.set_content("See forwarded email")
    outer.add_attachment(inner, filename="original.eml")  # how mail clients attach an email: message/rfc822
    (tmp_path / "fw.eml").write_bytes(outer.as_bytes())
    found = discover(tmp_path, "main")
    assert [d.display_name for d in found.documents] == ["fw.eml › original.eml › INV-9.pdf"]


def test_unsupported_attachment_listed(tmp_path):
    eml = make_eml([("old.doc", b"\xd0\xcf\x11\xe0junk", "application/msword")])
    (tmp_path / "e.eml").write_bytes(eml)
    found = discover(tmp_path, "main")
    assert found.skipped[0].status == "excluded"
    assert found.skipped[0].source.display_name == "e.eml › old.doc"
    assert found.documents[0].kind == "email_body"  # nothing usable attached, so the body is read


@pytest.mark.skipif(not (FIXTURES / "sample.msg").exists(), reason="No .msg fixture (made with Outlook, optional)")
def test_msg_fixture(tmp_path):
    found = discover(FIXTURES, "main")
    assert any(d.display_name.startswith("sample.msg") for d in found.documents)


# ---------------------------------------------------------------- discovery


def test_discovery_rules(tmp_path):
    (tmp_path / "b.pdf").write_bytes(make_pdf())
    (tmp_path / "a.png").write_bytes(big_photo((100, 100), "PNG"))
    (tmp_path / "old.doc").write_bytes(b"x")
    (tmp_path / "~$lock.docx").write_bytes(b"x")
    (tmp_path / ".hidden.pdf").write_bytes(b"x")
    (tmp_path / "Thumbs.db").write_bytes(b"x")
    sub = tmp_path / "sub"
    sub.mkdir()
    (sub / "c.pdf").write_bytes(make_pdf(text="other"))
    (sub / "copy.pdf").write_bytes(make_pdf(text="other"))  # identical bytes to c.pdf
    found = discover(tmp_path, "cutoff")
    assert [d.display_name for d in found.documents] == ["a.png", "b.pdf", "sub/c.pdf", "sub/copy.pdf"]
    assert found.files_found == 5
    assert found.skipped[0].reasons == ["Unsupported file type .doc. Save as .docx or PDF and re-run."]
    ids = [d.doc_id for d in found.documents]
    assert len(set(ids)) == 4
    assert ids[3] == ids[2] + "-2"
    assert all(d.test == "cutoff" for d in found.documents)


def test_limit(tmp_path):
    for i in range(5):
        (tmp_path / f"{i}.pdf").write_bytes(make_pdf(text=str(i)))
    assert [d.display_name for d in discover(tmp_path, "main", limit=2).documents] == ["0.pdf", "1.pdf"]
